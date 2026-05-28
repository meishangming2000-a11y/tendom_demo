from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\tendon_project")
HAND_ROOT = ROOT / "simulations" / "models" / "hand_stage1_export"
OUT_DIR = ROOT / "output" / "doc"
ASSET_DIR = OUT_DIR / "hand_stage1_training_literature_review_20260522_assets"
DOCX_PATH = OUT_DIR / "hand_stage1_training_literature_review_20260522.docx"
MD_PATH = OUT_DIR / "hand_stage1_training_literature_review_20260522_summary.md"


COLORS = {
    "ink": "#1f2933",
    "muted": "#52616b",
    "blue": "#2563eb",
    "teal": "#0f766e",
    "green": "#16a34a",
    "amber": "#d97706",
    "red": "#dc2626",
    "purple": "#7c3aed",
    "gray": "#e5e7eb",
    "light_blue": "#dbeafe",
    "light_green": "#dcfce7",
    "light_amber": "#fef3c7",
    "light_red": "#fee2e2",
}


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def set_cn_fonts() -> None:
    plt.rcParams["font.sans-serif"] = [
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def savefig(path: Path) -> Path:
    plt.tight_layout()
    plt.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def diagram_readiness_ladder() -> Path:
    steps = [
        ("CAD/URDF\nexport3", 0.9, "done"),
        ("clean visual\nSTL", 0.9, "done"),
        ("collision proxy\n+ actuators", 0.75, "done"),
        ("pinned-ball\nwrap", 0.70, "done"),
        ("zero-g\nfree release", 0.35, "risk"),
        ("gravity\nstable grasp", 0.18, "risk"),
        ("training\nbaseline", 0.25, "plan"),
    ]
    fig, ax = plt.subplots(figsize=(12.5, 3.4))
    ax.axis("off")
    x = np.arange(len(steps))
    y = np.zeros_like(x)
    ax.plot(x, y, color=COLORS["gray"], linewidth=10, solid_capstyle="round")
    for i, (label, score, status) in enumerate(steps):
        color = {
            "done": COLORS["green"],
            "risk": COLORS["amber"] if score > 0.25 else COLORS["red"],
            "plan": COLORS["blue"],
        }[status]
        ax.scatter(i, 0, s=900 * (0.7 + score), color=color, edgecolor="white", linewidth=2.5, zorder=3)
        ax.text(i, 0.34, label, ha="center", va="bottom", fontsize=11, color=COLORS["ink"], weight="bold")
        ax.text(i, -0.36, f"{int(score * 100)}%", ha="center", va="top", fontsize=10, color=COLORS["muted"])
    ax.text(
        0,
        -0.78,
        "Stage1/export3 已经能支持诊断式 scripted task；训练应从接口和低风险数据管线开始，而不是直接宣称 free-object grasp。",
        ha="left",
        va="top",
        fontsize=11,
        color=COLORS["ink"],
    )
    ax.set_xlim(-0.55, len(steps) - 0.45)
    ax.set_ylim(-1.05, 0.95)
    return savefig(ASSET_DIR / "fig01_readiness_ladder.png")


def diagram_method_taxonomy() -> Path:
    fig, ax = plt.subplots(figsize=(12.5, 6.2))
    ax.axis("off")
    main = [
        ("Scripted expert\n数据生成", COLORS["light_green"], COLORS["green"]),
        ("BC / DAgger\n模仿与纠错", COLORS["light_green"], COLORS["green"]),
        ("DAPG / PPO\nfrom demos", COLORS["light_amber"], COLORS["amber"]),
        ("zero-g / gravity\nfree-object eval", COLORS["light_blue"], COLORS["blue"]),
    ]
    x0, y0, w, h, gap = 0.07, 0.62, 0.18, 0.15, 0.055
    centers = []
    for i, (label, fc, ec) in enumerate(main):
        x = x0 + i * (w + gap)
        centers.append((x + w / 2, y0 + h / 2))
        ax.add_patch(plt.Rectangle((x, y0), w, h, fc=fc, ec=ec, lw=2))
        ax.text(x + w / 2, y0 + h / 2, label, ha="center", va="center", fontsize=11, color=COLORS["ink"], weight="bold")
        if i > 0:
            ax.annotate("", xy=(x - 0.006, y0 + h / 2), xytext=(x - gap + 0.006, y0 + h / 2), arrowprops=dict(arrowstyle="->", lw=1.8, color=COLORS["muted"]))

    supports = [
        ("Analytical / MPC\n安全限幅与姿态先验", 0.08, 0.32, COLORS["light_blue"], COLORS["blue"]),
        ("Diffusion / ACT / DP3\n第二阶段 sequence policy", 0.34, 0.32, "#ede9fe", COLORS["purple"]),
        ("DexGraspNet / UniDexGrasp\n多物体 proposal", 0.60, 0.32, COLORS["light_blue"], COLORS["blue"]),
        ("DexGen / VLA\n长期技能库与高层规划", 0.34, 0.10, COLORS["light_red"], COLORS["red"]),
    ]
    for label, x, y, fc, ec in supports:
        ax.add_patch(plt.Rectangle((x, y), 0.25, 0.13, fc=fc, ec=ec, lw=1.8))
        ax.text(x + 0.125, y + 0.065, label, ha="center", va="center", fontsize=10.5, color=COLORS["ink"], weight="bold")
    ax.annotate("", xy=(0.465, 0.62), xytext=(0.465, 0.45), arrowprops=dict(arrowstyle="->", lw=1.4, color=COLORS["purple"]))
    ax.annotate("", xy=(0.58, 0.62), xytext=(0.725, 0.45), arrowprops=dict(arrowstyle="->", lw=1.4, color=COLORS["blue"]))
    ax.text(0.06, 0.90, "机械手训练方法谱系：先形成可靠数据闭环，再升级模型复杂度", fontsize=14, weight="bold", color=COLORS["ink"])
    ax.text(0.06, 0.02, "对 hand_stage1：先 Scripted -> BC/DAgger -> Demo-augmented RL，再考虑 Diffusion/VLA。", fontsize=12, weight="bold", color=COLORS["ink"])
    return savefig(ASSET_DIR / "fig02_method_taxonomy.png")


def diagram_fit_heatmap() -> Path:
    methods = [
        "Scripted data + BC",
        "Interactive BC / DAgger",
        "Diffusion Policy",
        "DAPG / RL from demos",
        "PPO/SAC from scratch",
        "DexGraspNet-style proposal",
        "DexGen / VLA",
    ]
    criteria = [
        "fits current sim",
        "sample efficient",
        "handles contacts",
        "needs vision data",
        "compute burden",
        "near-term value",
    ]
    # Higher is better except compute burden and needs vision data, where lower need is better for current project.
    vals = np.array(
        [
            [5, 5, 2, 5, 5, 5],
            [4, 5, 3, 4, 4, 5],
            [3, 4, 3, 2, 3, 3],
            [3, 3, 5, 5, 2, 4],
            [2, 1, 5, 5, 1, 2],
            [2, 3, 3, 2, 3, 3],
            [1, 2, 4, 1, 1, 2],
        ],
        dtype=float,
    )
    fig, ax = plt.subplots(figsize=(12, 5.8))
    im = ax.imshow(vals, cmap="YlGnBu", vmin=1, vmax=5)
    ax.set_xticks(np.arange(len(criteria)), labels=criteria, rotation=25, ha="right", fontsize=10)
    ax.set_yticks(np.arange(len(methods)), labels=methods, fontsize=10)
    for i in range(vals.shape[0]):
        for j in range(vals.shape[1]):
            ax.text(j, i, int(vals[i, j]), ha="center", va="center", color="black", fontsize=10, weight="bold")
    ax.set_title("方法适配度矩阵：当前 hand_stage1/export3 的近中期选择", fontsize=14, weight="bold", pad=16)
    cbar = fig.colorbar(im, ax=ax, fraction=0.028, pad=0.02)
    cbar.set_label("1 = weak / costly, 5 = strong / practical", fontsize=9)
    return savefig(ASSET_DIR / "fig03_fit_heatmap.png")


def diagram_training_pipeline() -> Path:
    fig, ax = plt.subplots(figsize=(12.5, 6.0))
    ax.axis("off")
    lanes = [
        ("A. Sim contract gate", "load scene, no overlap, contacts, success metrics", COLORS["blue"], 0.82),
        ("B. Scripted expert", "pinned -> zero-g -> gravity, sweep + labels", COLORS["green"], 0.66),
        ("C. Imitation learning", "BC/DAgger first, sequence model second", COLORS["teal"], 0.50),
        ("D. RL fine-tune", "DAPG-style demo start, PPO/SAC only after reward is stable", COLORS["amber"], 0.34),
        ("E. Generalization", "domain randomization, object set, point cloud/diffusion", COLORS["purple"], 0.18),
    ]
    for idx, (title, desc, color, y) in enumerate(lanes):
        ax.add_patch(plt.Rectangle((0.06, y - 0.055), 0.86, 0.11, fc="white", ec=color, lw=2))
        ax.text(0.08, y + 0.018, title, ha="left", va="center", fontsize=12, weight="bold", color=color)
        ax.text(0.36, y + 0.018, desc, ha="left", va="center", fontsize=10.5, color=COLORS["ink"])
        if idx < len(lanes) - 1:
            ax.annotate("", xy=(0.49, lanes[idx + 1][3] + 0.065), xytext=(0.49, y - 0.065), arrowprops=dict(arrowstyle="->", lw=1.8, color=COLORS["muted"]))
    ax.text(0.06, 0.94, "推荐训练路线：先让数据和指标可靠，再让模型复杂", fontsize=15, weight="bold", color=COLORS["ink"])
    ax.text(0.06, 0.04, "Stop condition: free-ball release retention and contact quality must pass before RL is treated as a benchmark baseline.", fontsize=10.5, color=COLORS["red"], weight="bold")
    return savefig(ASSET_DIR / "fig04_training_pipeline.png")


def diagram_task_contract() -> Path:
    fig, ax = plt.subplots(figsize=(12.3, 5.2))
    ax.axis("off")
    boxes = [
        ("Observation", "qpos/qvel\nobject pose\nfingertip distances\ncontact summary", 0.08, 0.56, COLORS["light_blue"], COLORS["blue"]),
        ("Policy", "MLP BC -> DAgger\nsequence BC/diffusion\nDAPG/PPO fine-tune", 0.39, 0.56, COLORS["light_green"], COLORS["green"]),
        ("Action", "21-D actuator targets\nposition control\nrate limits", 0.70, 0.56, COLORS["light_amber"], COLORS["amber"]),
        ("Evaluator", "no initial overlap\ncontact support\npenetration cap\nrelease retention", 0.39, 0.20, COLORS["light_red"], COLORS["red"]),
    ]
    for title, body, x, y, fc, ec in boxes:
        ax.add_patch(plt.Rectangle((x, y), 0.23, 0.22, fc=fc, ec=ec, lw=2))
        ax.text(x + 0.115, y + 0.17, title, ha="center", va="center", fontsize=12, weight="bold", color=ec)
        ax.text(x + 0.115, y + 0.075, body, ha="center", va="center", fontsize=10, color=COLORS["ink"])
    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="->", lw=1.7, color=COLORS["muted"]))
    arrow(0.31, 0.67, 0.39, 0.67)
    arrow(0.62, 0.67, 0.70, 0.67)
    arrow(0.815, 0.56, 0.515, 0.42)
    arrow(0.39, 0.31, 0.19, 0.56)
    ax.text(0.08, 0.90, "训练前必须冻结的 task contract", fontsize=15, weight="bold", color=COLORS["ink"])
    ax.text(0.08, 0.08, "这张图对应仓库现有思路：task-level official metrics 与 backend debug metrics 分离。", fontsize=10.5, color=COLORS["muted"])
    return savefig(ASSET_DIR / "fig05_task_contract.png")


def make_diagrams() -> list[Path]:
    set_cn_fonts()
    return [
        diagram_readiness_ladder(),
        diagram_method_taxonomy(),
        diagram_fit_heatmap(),
        diagram_training_pipeline(),
        diagram_task_contract(),
    ]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True)
        set_cell_shading(hdr[i], "E8EEF7")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    doc.add_paragraph(text, style=style)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(82, 97, 107)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10)
    for style_name, size, color in [
        ("Title", 22, "1F2933"),
        ("Heading 1", 16, "1F2933"),
        ("Heading 2", 13, "2563EB"),
        ("Heading 3", 11, "0F766E"),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def reference_rows() -> list[list[str]]:
    return [
        ["R1", "Rajeswaran et al., Learning Complex Dexterous Manipulation with Deep RL and Demonstrations, RSS 2018", "DAPG; demonstrations reduce sample complexity for 24-DoF Adroit hand", "https://arxiv.org/abs/1709.10087"],
        ["R2", "OpenAI et al., Learning Dexterous In-Hand Manipulation, 2018/2019", "Pure sim RL with heavy domain randomization on Shadow Hand", "https://arxiv.org/abs/1808.00177"],
        ["R3", "OpenAI et al., Solving Rubik's Cube with a Robot Hand, 2019", "Automatic Domain Randomization for sim-to-real dexterity", "https://arxiv.org/abs/1910.07113"],
        ["R4", "Handa et al., DeXtreme, ICRA 2023", "Isaac Gym RL, vision/state policies, robust sim-to-real dexterous reorientation", "https://arxiv.org/abs/2210.13702"],
        ["R5", "Qin et al., DexPoint, CoRL 2022", "Point-cloud RL for generalizable sim-to-real dexterous manipulation", "https://arxiv.org/abs/2211.09423"],
        ["R6", "Xu et al., UniDexGrasp, CVPR 2023", "Grasp proposal + goal-conditioned execution for many objects", "https://arxiv.org/abs/2303.00938"],
        ["R7", "Chi et al., Diffusion Policy, RSS 2023 / journal 2024", "Action diffusion handles multimodal high-dimensional actions", "https://arxiv.org/abs/2303.04137"],
        ["R8", "Ze et al., 3D Diffusion Policy, RSS 2024", "Point-cloud-conditioned diffusion policy; strong few-demo performance", "https://arxiv.org/abs/2403.03954"],
        ["R9", "Wang et al., DexCap, 2024", "Portable human hand mocap + DexIL for dexterous imitation", "https://arxiv.org/abs/2403.07788"],
        ["R10", "Zhang et al., DexGraspNet 2.0, 2024", "427M synthetic dexterous grasps; local-geometry diffusion proposal", "https://arxiv.org/abs/2410.23004"],
        ["R11", "Jiang et al., DexMimicGen, ICRA 2025", "Turns a small number of demos into large simulated dexterous datasets", "https://arxiv.org/abs/2410.24185"],
        ["R12", "Yin et al., DexterityGen / DexGen, 2025", "RL pretrains primitive dexterity, human prompts compose skills", "https://arxiv.org/abs/2502.04307"],
        ["R13", "Lin et al., Sim-to-Real RL for Vision-Based Dexterous Manipulation on Humanoids, CoRL 2025", "Real-to-sim tuning, generalized reward, distillation for contact-rich tasks", "https://arxiv.org/abs/2502.20396"],
        ["R14", "Zhong et al., DexGraspVLA, 2025", "VLM planner + diffusion low-level controller for general dexterous grasping", "https://arxiv.org/abs/2502.20900"],
        ["R15", "Xu et al., DexUMI, 2025", "Human hand as universal interface with exoskeleton and visual adaptation", "https://arxiv.org/abs/2505.21864"],
        ["R16", "Weinberg et al., Survey of learning-based approaches for robotic in-hand manipulation, Frontiers 2024", "Taxonomy of model-based, RL, and IL for in-hand manipulation", "https://doi.org/10.3389/frobt.2024.1455431"],
        ["R17", "Welte and Rayyes, Interactive imitation learning survey, Frontiers 2025", "IIL for high-dimensional contact-rich dexterous manipulation", "https://doi.org/10.3389/frobt.2025.1682437"],
    ]


def create_docx(diagrams: list[Path]) -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("hand_stage1/export3 机械手训练路线文献综述")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string("1F2933")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("从 Shadow-style scripted scaffold 到可训练策略的分阶段路线")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("52616B")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Generated: {date.today().isoformat()} | Repository: D:\\tendon_project").italic = True

    doc.add_paragraph()
    add_figure(doc, diagrams[0], "图 1. 当前 hand_stage1/export3 训练就绪度阶梯。", width=6.7)

    doc.add_heading("0. 执行摘要", level=1)
    doc.add_paragraph(
        "结论先说清楚：我们可以开始做 training 相关的工程准备和低风险训练实验，但不应直接开始把当前模型当作稳定自由抓取基准来训练。"
        "最合理的起点不是纯 RL，而是把 export3 已跑通的 Shadow-style scripted scaffold 变成数据生成器、任务合约和评估器，先训练 state-based BC/DAgger，"
        "再用 demonstration-augmented RL 做 free-ball retention 的细化。"
    )
    add_bullet(doc, "当前 export3 已具备：可加载 MuJoCo scene、clean visual STL、primitive collision proxy、21 个 position actuators、5 个 fingertip sites、pinned-ball wrap smoke test。")
    add_bullet(doc, "当前最大缺口：free-object retention 没过；zero-g release 仍会漂移，gravity release 会掉开；collision proxy、contact/friction、thumb axis/limit 仍需确认。")
    add_bullet(doc, "推荐近期路线：Scripted expert dataset -> BC baseline -> DAgger/interactive correction -> DAPG/PPO fine-tune。")
    add_bullet(doc, "Diffusion Policy/DP3 值得作为第二阶段方法，但应等任务数据量、视觉/点云输入和多模态动作分布稳定后再上。")
    add_bullet(doc, "DexGen/VLA/DexUMI/DexGraspNet 2.0 是重要方向，但对我们现在更像路线参考或中长期模块，而不是下周的主训练 baseline。")

    doc.add_heading("1. 我们已经做过的工作回顾", level=1)
    doc.add_paragraph(
        "本节把 hand_stage1/export3 的现状压缩成训练视角的事实，而不是 CAD 或 mesh 调试流水账。"
    )
    rows = [
        ["URDF / tree", "export3 thumb chain 已变成 2-DoF CMC: thumb_cmc_abd_joint + thumb_cmc_flex_joint。wrist_2_joint 当前为 fixed，需要人工确认。"],
        ["visual mesh", "export3 STL 24 个，hash 不重复，visual-only 使用；不再使用旧 suspicious mesh。"],
        ["MuJoCo model", "scene_ball_export3.xml 可加载；26 bodies including world, 22 joints including ball freejoint, 21 position actuators, 55 geoms, 5 sites, 24 meshes。"],
        ["collision", "当前是 simplified primitive proxy，不使用 STL collision。可 smoke test，但不是最终 contact fidelity。"],
        ["scripted grasp", "pinned-ball wrap 通过；best hold: 5 contacts, max penetration about 0.00233 m, mean four-tip distance about 0.03992 m, thumb-ball about 0.07693 m。"],
        ["release test", "zero-g release retained 未过；gravity release retained 未过。当前最多是 pinned object scripted wrap，不是 stable free-object grasp。"],
        ["official repo baseline", "仓库官方主线仍是 pre_grasp expert -> dataset -> BC train -> eval -> report。stable_grasp 是结构化 skeleton，不是 promoted training baseline。"],
    ]
    add_table(doc, ["项目层面", "训练相关事实"], rows, widths=[4.0, 12.0])

    doc.add_heading("2. 文献方法地图", level=1)
    add_figure(doc, diagrams[1], "图 2. 机械手训练方法谱系以及推荐先后顺序。", width=6.7)
    doc.add_paragraph(
        "近几年机械手训练大体可以分成七类：传统/模型驱动控制、scripted expert、行为克隆与 DAgger、diffusion/sequence imitation、model-free RL、"
        "RL + demonstration、以及大规模生成数据/基础模型路线。它们不是互斥的；真正成功的 dexterous 系统往往是组合式。"
    )

    method_rows = [
        ["Analytical / MPC / trajectory optimization", "需要精确动力学和接触模型；可解释但对高维多接触很脆弱。", "适合做安全约束和低级限幅，不适合作为当前主训练路线。"],
        ["Scripted expert", "工程可控、数据便宜、容易诊断失败；覆盖范围有限。", "当前最适合。把 scripted sweep 变成 dataset generator。"],
        ["BC / DAgger / IIL", "样本效率高，能继承 scripted/human corrections；BC 有 covariate shift。", "近期主线。先 state-based，再加 interactive correction。"],
        ["Diffusion Policy / DP3", "擅长高维、多模态、时序动作；训练稳定，但需要较好的 demos 和输入表征。", "第二阶段。适合 scripted demos 变多之后。"],
        ["DAPG / RL from demonstrations", "兼顾 demo prior 和 RL contact discovery；比纯 RL 更现实。", "free-ball retention 阶段首选 RL 路线。"],
        ["PPO/SAC from scratch", "能发现复杂接触策略，但奖励和采样压力大。", "不建议现在直接上；可作为后期 ablation。"],
        ["DexGraspNet/UniDexGrasp-style proposal", "学习 grasp pose proposal + execution，适合多物体泛化。", "中期用于多物体目标姿态生成，不是当前第一 baseline。"],
        ["DexGen/VLA/foundation controller", "趋势很强，但需要大规模数据、真实/仿真多任务和接口成熟。", "长期方向；可借鉴分层控制，不宜现在押主线。"],
    ]
    add_table(doc, ["方法族", "核心逻辑", "对本项目的判断"], method_rows, widths=[4.0, 6.0, 6.0])

    doc.add_heading("3. 重点文献综述", level=1)

    doc.add_heading("3.1 RL 和 sim-to-real 路线", level=2)
    doc.add_paragraph(
        "OpenAI Dactyl 和 Rubik's Cube 工作证明了大规模仿真、domain randomization 和 recurrent policy 可以把 Shadow Hand 的高维控制从仿真带到真实系统。"
        "但这类路线对仿真质量、随机化、算力、reward design 和真实平台稳定性要求很高。对 hand_stage1 来说，它的价值在于告诉我们未来要做 domain randomization，"
        "而不是现在直接复制大规模 RL。"
    )
    add_bullet(doc, "R2: Dactyl 使用仿真训练和物理参数/视觉随机化，实现物体 reorientation 的 sim-to-real。")
    add_bullet(doc, "R3: Rubik's Cube 引入 automatic domain randomization，让环境难度逐步扩展。")
    add_bullet(doc, "R4: DeXtreme 使用 Isaac Gym 和 Allegro Hand，强调 GPU 并行、鲁棒 pose estimator 和 sim-to-real pipeline。")
    add_bullet(doc, "R13: 2025 humanoid dexterous RL 使用 real-to-sim tuning、general reward、divide-and-conquer distillation，说明近期 RL 越来越强调工程配方而不是单一算法。")

    doc.add_heading("3.2 Demonstration + RL: DAPG 是当前最重要的近中期参考", level=2)
    doc.add_paragraph(
        "Rajeswaran et al. 的 DAPG 对我们非常相关：高维 Adroit hand 任务可以从少量 demonstration 起步，再用 policy gradient fine-tune。"
        "这和我们现在的局面吻合：我们已经有 scripted expert，但自由抓取没过；纯 RL 太难，纯 BC 又可能在 release 时分布外崩掉。"
    )
    add_bullet(doc, "近期建议把 scripted pose/position-control rollout 存为 demos，先训练 BC，再做 demonstration-augmented PPO/DAPG-style fine-tune。")
    add_bullet(doc, "reward 不应一开始复杂化；先围绕 no-overlap、contact support、penetration cap、fingertip-ball distance、release retention。")

    doc.add_heading("3.3 Imitation learning、DAgger 和 interactive imitation", level=2)
    doc.add_paragraph(
        "2025 的 interactive imitation survey 强调一个问题：dexterous manipulation 的高维接触任务中，纯离线 BC 很容易出现 covariate shift，"
        "而人类纠错/on-policy intervention 能显著改善安全性和样本效率。我们短期没有真实手 teleop，但可以做仿真中的人工/脚本纠错版 DAgger。"
    )
    add_bullet(doc, "第一版：scripted expert -> BC。")
    add_bullet(doc, "第二版：BC rollout 失败处由 scripted solver 或人工选择修正目标，聚合回 dataset。")
    add_bullet(doc, "第三版：如果以后有手套/视觉/遥操作，参考 DexCap/DexUMI 做 human-to-robot data collection。")

    doc.add_heading("3.4 Diffusion Policy / DP3: 很适合后续，但不应抢第一棒", level=2)
    doc.add_paragraph(
        "Diffusion Policy 把动作序列看成条件生成问题，优点是能处理多模态动作、高维 action 和训练稳定性。DP3 进一步强调 sparse point cloud 的 3D 表征，"
        "在少量 demonstration 上表现强。对我们来说，diffusion 是第二阶段很有吸引力的模型：当 scripted/DAgger 数据有几千到几万条 rollout，并且 observation schema 固定后，"
        "它可能比 MLP BC 更能表达不同手指闭合时序。"
    )
    add_bullet(doc, "现在不建议直接上视觉 diffusion；先用 state/action 序列版或轻量 Transformer/ACT 做 ablation。")
    add_bullet(doc, "一旦加点云或多物体，DP3 比 2D image policy 更贴近 contact geometry。")

    doc.add_heading("3.5 大规模 grasp proposal 和 foundation controller", level=2)
    doc.add_paragraph(
        "UniDexGrasp、DexGraspNet 2.0、DexGraspVLA 和 DexGen 代表了更宏观的趋势：把抓取拆成 proposal、low-level execution、foundation planner/controller，"
        "并用大规模合成数据或 foundation model 增强泛化。这些工作值得我们在接口设计上提前兼容，但它们需要更稳定的手模型、视觉/点云、对象库和评估体系。"
    )
    add_bullet(doc, "DexGraspNet 2.0 的 local geometry diffusion proposal 可作为中期多物体抓取姿态生成器参考。")
    add_bullet(doc, "DexGen 的 primitive controller 思路适合未来把 close、roll、translate、reorient 做成技能库。")
    add_bullet(doc, "DexGraspVLA/VLA 更适合长期高层任务规划，不适合当前低级接触还未稳定时直接导入。")

    doc.add_heading("4. 对当前 hand_stage1/export3 的适配度判断", level=1)
    add_figure(doc, diagrams[2], "图 3. 方法适配度矩阵。分数是面向当前 hand_stage1/export3 状态的工程判断。", width=6.7)
    doc.add_paragraph(
        "从适配度看，当前最应推进的不是最复杂模型，而是最能把现有 scripted scaffold 转化成可复现实验资产的路线。"
    )
    recommendation_rows = [
        ["1", "state-based BC", "马上做", "已有 21-D actuator、object pose、contact/fingertip metrics；和仓库官方 pre_grasp BC 主线一致。", "训练 pinned/pre-release close，不把它叫 stable free grasp。"],
        ["2", "DAgger / interactive correction", "马上设计，随后做", "解决 BC rollout 分布偏移；失败点可由 scripted solver 纠正。", "需要统一 dataset schema 和 failure labels。"],
        ["3", "DAPG-style RL fine-tune", "collision proxy 过门后做", "最适合从 demos 推进到 zero-g/free-ball retention。", "需要 reward、domain randomization 和并行 rollout。"],
        ["4", "Diffusion Policy / ACT", "中期做 ablation", "动作序列多模态，适合手指闭合节奏。", "数据量和 observation schema 稳定后再上。"],
        ["5", "DexGraspNet/UniDexGrasp-style proposal", "中长期", "用于多物体、多姿态 grasp target proposal。", "需要对象库、点云和稳定 execution policy。"],
        ["6", "VLA/DexGen foundation route", "长期路线", "适合任务级泛化和技能组合。", "当前模型还没过 free grasp，不宜作为主线。"],
    ]
    add_table(doc, ["优先级", "模型/路线", "时间点", "为什么适合", "注意事项"], recommendation_rows, widths=[1.2, 3.3, 2.5, 5.3, 4.0])

    doc.add_heading("5. 推荐训练路线", level=1)
    add_figure(doc, diagrams[3], "图 4. 推荐分阶段训练路线。", width=6.7)

    doc.add_heading("5.1 Phase 0: 训练前必须补齐的 contract", level=2)
    add_figure(doc, diagrams[4], "图 5. 训练前应冻结的 observation-action-evaluator 合约。", width=6.7)
    add_bullet(doc, "Observation: qpos/qvel, object pose/velocity, fingertip site positions, fingertip-ball distances, contact count, penetration, previous action。")
    add_bullet(doc, "Action: 21-D position actuator target 或 target delta；必须限幅、限速、记录 clipped action。")
    add_bullet(doc, "Success: no initial overlap, sufficient contact support, penetration below threshold, retention after release, no floor contact。")
    add_bullet(doc, "Episode labels: success/failure/timeout/initial_overlap/contact_lost/object_dropped/thumb_far。")

    doc.add_heading("5.2 Phase 1: Scripted dataset + BC baseline", level=2)
    doc.add_paragraph(
        "把已有 `run_export3_shadow_style_scripted_task.py` 发展成 collector：对球位、手指尺度、拇指 pose、速度、摩擦、碰撞半径做 sweep，保存每步 observation/action。"
        "第一版训练目标不是自由抓取，而是 imitation of good pinned/pre-release trajectories。"
    )
    add_bullet(doc, "Dataset v0: pinned_wrap_success trajectories + selected partial/failure trajectories for negative labels。")
    add_bullet(doc, "Model v0: MLP BC, input state vector, output 21-D actuator target delta。")
    add_bullet(doc, "Metrics: imitation MSE, rollout contact score, fingertip distance, penetration, release drift。")

    doc.add_heading("5.3 Phase 2: DAgger / interactive correction", level=2)
    doc.add_paragraph(
        "让 BC policy 自己 rollout，遇到 drift/contact loss/thumb far 时调用 scripted correction 或人工选择 correction pose，聚合回数据集。"
        "这一步比直接 RL 更稳，因为它先修正 state distribution mismatch。"
    )

    doc.add_heading("5.4 Phase 3: Demo-augmented RL for zero-g release", level=2)
    doc.add_paragraph(
        "当 collision proxy 和 task metric 稳定后，做 DAPG-style 或 PPO-with-BC-regularization。第一目标是 zero-g release retention，"
        "而不是一上来做 gravity lift。reward 要尽量短：contact support、object retained、low penetration、thumb/index opposition、action smoothness。"
    )

    doc.add_heading("5.5 Phase 4: Gravity, object set, domain randomization", level=2)
    doc.add_paragraph(
        "只有 zero-g free release 过门后，才进入 gravity。此时加入 friction/mass/object radius/randomized pose。参考 Dactyl/ADR/DeXtreme，"
        "但保持小规模、可解释的 randomization schedule。"
    )

    doc.add_heading("6. 训练任务定义建议", level=1)
    task_rows = [
        ["hand_stage1_pre_grasp", "把掌心和手指移动到球附近，无初始穿透", "palm/object relation, qpos/qvel", "21-D position target", "no overlap + distance/orientation tolerance"],
        ["hand_stage1_pinned_wrap", "在 pinned 球上形成多指接触和包络", "qpos/qvel, ball pose, fingertip distances", "21-D position target", ">=3-5 contacts, penetration cap, thumb not too far"],
        ["hand_stage1_zero_g_release", "闭合后释放球，zero-g 下保持在掌内", "state + contact + ball velocity", "21-D target/delta", "release drift below threshold for N steps"],
        ["hand_stage1_gravity_hold", "重力下保持小球不掉落", "state + contact + object height", "21-D target/delta", "no floor contact, retained workspace, stable contacts"],
        ["hand_stage1_lift_and_hold", "未来任务，形成 grasp 后移动/抬起", "需要 wrist/arm or platform", "hand + arm action", "not ready in current hand-only scene"],
    ]
    add_table(doc, ["任务名", "目标", "Observation", "Action", "Success"], task_rows, widths=[3.4, 4.0, 4.0, 3.0, 4.0])

    doc.add_heading("7. 风险、阻塞和 Go/No-Go", level=1)
    risk_rows = [
        ["collision proxy", "MAJOR", "当前球释放失败很可能和接触几何/摩擦/指尖 proxy 有关。", "先调 proxy，再训练 RL。"],
        ["thumb semantics", "MAJOR", "export3 改善明显，但 opposition 仍不是最终 Shadow-like clamp。", "SolidWorks 里确认 thumb_cmc_abd/flex axis 和 limits。"],
        ["free-object retention", "BLOCKER for stable_grasp training", "release 后球离手。", "先 zero-g release pass，再 gravity。"],
        ["reward hacking", "MAJOR", "RL 可能通过穿透或夹爆方式刷分。", "penetration cap 和 visual audit 必须进 official eval。"],
        ["data bias", "MAJOR", "scripted demos 覆盖范围窄。", "DAgger + failure-state sampling。"],
        ["compute", "MINOR/MAJOR", "MuJoCo CPU sweep 可以起步，大规模 RL 可能慢。", "先小模型；需要时考虑 MJX/Isaac Lab 迁移。"],
    ]
    add_table(doc, ["风险", "级别", "说明", "建议"], risk_rows, widths=[3.4, 2.0, 6.0, 5.0])

    doc.add_heading("8. 具体下一步实验计划", level=1)
    add_bullet(doc, "E1: 写 `collect_export3_scripted_dataset.py`，复用 Shadow-style sweep，保存 NPZ + JSON report。")
    add_bullet(doc, "E2: 写 `train_export3_bc.py`，先训练 state-only MLP BC；输出 `bc_hand_stage1_export3_pinned_wrap_v0.pth`。")
    add_bullet(doc, "E3: 写 `eval_export3_bc.py`，评估 pinned、zero-g release、gravity release 三档。")
    add_bullet(doc, "E4: 增加 DAgger loop：BC 失败处调用 scripted correction，生成 `expert+correction` 数据集。")
    add_bullet(doc, "E5: collision proxy sweep：指尖 sphere 半径、palm proxy、friction、球半径/质量。")
    add_bullet(doc, "E6: DAPG/PPO fine-tune，只在 zero-g release gate 达标后启动。")
    add_bullet(doc, "E7: 中期 ablation：sequence BC / ACT / Diffusion Policy，比较 rollout 成功率而不是只看 loss。")

    doc.add_heading("9. 推荐决策", level=1)
    doc.add_paragraph(
        "我建议从现在开始 training，但把它命名为 `diagnostic training line` 或 `stage1.1 training scaffold`，不要叫正式 stable_grasp baseline。"
        "第一阶段目标是训练一个能复现 scripted pinned-wrap 的策略，并用它暴露 task contract、dataset schema、eval report、failure labels。"
        "第二阶段才是 zero-g free release；第三阶段再做 gravity。"
    )
    add_bullet(doc, "立刻可做：BC/DAgger 数据管线和 pinned-wrap imitation。")
    add_bullet(doc, "谨慎推进：DAPG/PPO fine-tune for zero-g release。")
    add_bullet(doc, "暂不主推：纯 RL from scratch、VLA/foundation controller、直接多物体泛化。")
    add_bullet(doc, "文档里的核心态度：可以训练，但训练应该先服务于工程诊断和接口稳定，而不是马上追求最终 dexterous benchmark。")

    doc.add_heading("10. 参考文献与来源", level=1)
    add_table(doc, ["ID", "文献", "与本文相关点", "URL"], reference_rows(), widths=[1.0, 6.0, 5.6, 5.0])

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Appendix A. 本项目建议的最小数据 schema", level=1)
    schema = {
        "episode": ["task_name", "model_name", "scene_xml", "seed", "mode", "ball_position", "success", "failure_reason"],
        "per_step_obs": ["qpos[21]", "qvel[21]", "ball_pos[3]", "ball_vel[3]", "tip_pos[5,3]", "tip_ball_dist[5]", "contact_count", "max_penetration"],
        "per_step_action": ["ctrl_target[21]", "ctrl_delta[21]", "clipped_flags[21]"],
        "labels": ["stage", "contact_support", "thumb_far", "object_retained", "floor_contact", "release_drift"],
    }
    for k, vals in schema.items():
        doc.add_heading(k, level=2)
        add_bullet(doc, ", ".join(vals))

    doc.save(DOCX_PATH)


def create_markdown_summary() -> None:
    lines = [
        "# hand_stage1/export3 训练路线文献综述摘要",
        "",
        f"- DOCX: `{DOCX_PATH}`",
        f"- Assets: `{ASSET_DIR}`",
        "- Recommendation: start training infrastructure now, but first as diagnostic BC/DAgger over scripted data, not as a promoted stable free-object grasp baseline.",
        "- Best current demonstrated capability: pinned-ball wrap/contact smoke test.",
        "- Main blocker: free-object release retention.",
        "",
        "## Core staged route",
        "",
        "1. Scripted expert dataset.",
        "2. State-based BC.",
        "3. DAgger / interactive correction.",
        "4. Demo-augmented RL for zero-g release.",
        "5. Gravity and domain randomization.",
        "6. Later: diffusion / point-cloud / VLA modules.",
    ]
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    ensure_dirs()
    diagrams = make_diagrams()
    create_docx(diagrams)
    create_markdown_summary()
    print(json.dumps({"docx": str(DOCX_PATH), "summary": str(MD_PATH), "assets": [str(p) for p in diagrams]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
